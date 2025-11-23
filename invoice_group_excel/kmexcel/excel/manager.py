from datetime import date, time, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import Cell
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, column_index_from_string
from openpyxl.worksheet.worksheet import Worksheet


from invoice_group_excel.kmexcel.excel.types.formula import FormulaDefinition, FormulaTemplate
from invoice_group_excel.kmexcel.excel.types.style import SheetStyle, apply_style, CellStyle


class ExcelManager:
    def __init__(self, filepath: str) -> None:
        self.filepath = filepath
        self._load_or_create()

    def save(self, filepath: str | None = None) -> None:
        if filepath:
            self.workbook.save(filepath)
        else:
            self.workbook.save(self.filepath)

    def get_sheet(self, name: str) -> Worksheet:
        return self.workbook[name]

    def get_sheet_names(self) -> list[str]:
        return self.workbook.sheetnames

    def remove_sheet(self, name: str) -> None:
        if name not in self.workbook.sheetnames:
            raise ValueError(f'Sheet "{name}" not found')
        del self.workbook[name]

    # ------------------------------------------------------------------------------------------------------------------
    # DANE
    # ------------------------------------------------------------------------------------------------------------------
    def read_sheet(self, sheet_name: str, exception_read: dict[str,list[Any]] | None = None) -> list[dict[str, Any]]:
        dict_keys = []
        list_of_dicts = []
        sheet = self.get_sheet(sheet_name)
        for i,row in enumerate(sheet.iter_rows()):
            if i == 0:
                dict_keys = [cell.value for cell in row]
            else:
                dict_data = dict(zip(dict_keys,  [cell.value for cell in row]))
                if not exception_read:
                    list_of_dicts.append(dict_data)
                else:
                    skip = False
                    for ex_k,ex_v in exception_read.items() :
                        if row[column_index_from_string(ex_k)] == ex_v:
                            skip = True
                    if skip:
                        continue
                    else:
                        list_of_dicts.append(dict_data)

        return list_of_dicts



    def add_sheet[T:dict](
            self,
            name: str,
            data: list[T],
            formulas: list[FormulaDefinition] | None = None,
            formula_templates: list[FormulaTemplate] | None = None
    ) -> None:
        if name in self.workbook.sheetnames:
            del self.workbook[name]

        ws = self.workbook.create_sheet(title=name)
        if not data:
            return

        headers = list(data[0].keys())
        ws.append(headers)

        # Wiersze z danymi
        for row in data:
            ws.append(list(row.values()))

        # Formuly bezposrednie
        if formulas:
            for formula in formulas:
                ws[formula['cell']].value = formula['formula']

        # Formuly szablonowe
        if formula_templates:
            last_row = ws.max_row
            for tmpl in formula_templates:
                col = tmpl['column']
                start = tmpl.get('start_row', 2)
                end = tmpl.get('end_row', last_row)
                template = tmpl['template']
                for row_idx in range(start, end + 1):
                    cell = ws[f'{col}{row_idx}']
                    cell.value = template.format(row=row_idx, col=col, last=last_row)

    def add_column[T: (bool | float | Decimal | str | date | time | timedelta)](
            self,
            sheet_name: str,
            values: list[T],
            col_letter: str = 'A',
            row_start: int = 1
    ) -> None:
        ws = self.get_sheet(sheet_name)
        col_index = column_index_from_string(col_letter)

        for idx, val in enumerate(values):
            ws.cell(row=row_start + idx, column=col_index).value = val

    def add_row[T: dict](
            self,
            sheet_name: str,
            data: T,
            row_index: int = 1,
            col_letter: str = 'A'
    ) -> None:
        ws = self.get_sheet(sheet_name)
        start_col_idx = column_index_from_string(col_letter)
        for offset, value in enumerate(data.values()):
            ws.cell(row=row_index, column=start_col_idx + offset).value = value

    def add_cell[T](self, sheet_name: str, cell_ref: str, value: T) -> None:
        ws = self.get_sheet(sheet_name)
        ws[cell_ref].value = value

    # ------------------------------------------------------------------------------------------------------------------
    # FORMAT DANYCH
    # ------------------------------------------------------------------------------------------------------------------

    def set_column_format(
            self,
            sheet_name: str,
            col_letter: str,
            number_format: str,
            start_row: int = 1,
            end_row: int | None = None) -> None:
        ws = self.get_sheet(sheet_name)
        end_row = end_row or ws.max_row
        col_index = column_index_from_string(col_letter)

        for row_idx in range(start_row, end_row + 1):
            cell = ws.cell(row=row_idx, column=col_index)
            cell.number_format = number_format

    # ------------------------------------------------------------------------------------------------------------------
    # STYLE
    # ------------------------------------------------------------------------------------------------------------------

    def style_all(self, sheet_name: str, style: CellStyle) -> None:
        ws = self.get_sheet(sheet_name)

        for row_idx in range(1, ws.max_row + 1):
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                if isinstance(cell, Cell):
                    apply_style(cell, style)

    def style_column(
            self,
            sheet_name: str,
            col_letter: str,
            style: CellStyle,
            start_row: int = 2,
            end_row: int | None = None
    ) -> None:
        ws = self.get_sheet(sheet_name)
        end_row = end_row or ws.max_row
        col_index = column_index_from_string(col_letter)

        for row in range(start_row, end_row + 1):
            cell = ws.cell(row=row, column=col_index)
            if isinstance(cell, Cell):
                apply_style(cell, style)

    def style_row(
            self,
            sheet_name: str,
            row_index: int,
            style: CellStyle,
            start_col_letter: str = 'A',
            end_col_letter: str | None = None
    ) -> None:
        ws = self.get_sheet(sheet_name)
        start_col = column_index_from_string(start_col_letter)
        end_col = column_index_from_string(end_col_letter) if end_col_letter else ws.max_column

        for col in range(start_col, end_col + 1):
            if row_index < 0:
                cell = ws.cell(row=ws.max_row + 1 + row_index, column=col)
            else:
                cell = ws.cell(row=row_index, column=col)
            if isinstance(cell, Cell):
                apply_style(cell, style)

    def style_cell(self, sheet_name: str, cell_ref: str, style: CellStyle) -> None:
        ws = self.get_sheet(sheet_name)
        cell = ws[cell_ref]
        apply_style(cell, style)

    def apply_cell_styles(self, sheet_name: str, headers: list[str], style: SheetStyle) -> None:
        ws = self.get_sheet(sheet_name)
        if style and 'header' in style:
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                if isinstance(cell, Cell):
                    apply_style(cell, style['header'])

        if style and 'row' in style:
            for row_idx in range(2, ws.max_row + 1):
                for col_idx in range(1, len(headers) + 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    if isinstance(cell, Cell):
                        apply_style(cell, style['row'])

    # ------------------------------------------------------------------------------------------------------------------
    # AUTODOPASOWANIE SZEROKOSCI W ZALEZNOSCI OD ZAWARTOSCI DANYCH W POSZCZEGOLNYCH KOLUMNACH
    # ------------------------------------------------------------------------------------------------------------------
    def autofit_column_widths(self, sheet_name: str, offset_dim: int = 3) -> None:

        ws: Worksheet = self.get_sheet(sheet_name)

        # Iteruj przez wszystkie kolumny w arkuszu od wiersza 1 do ostatniego wiersza z danymi
        for col_cells in ws.iter_cols(min_row=1, max_row=ws.max_row):
            # Pobierz numer kolumny (np. 1 dla A, 2 dla B, ...)
            col_idx = col_cells[0].column
            max_length = 0

            # Iteruj przez wszystkie komorki w biezacej kolumnie
            for cell in col_cells:
                value = cell.value
                if value is not None:
                    try:
                        max_length = max(max_length, len(str(value)))
                    except Exception:
                        continue

            # Zmien numer kolumny na litere
            if col_idx:
                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max_length + offset_dim

    # ------------------------------------------------------------------------------------------------------------------
    # GENEROWANIE WYKRESOW
    # ------------------------------------------------------------------------------------------------------------------

    def add_new_chart(
            self,
            chart_type: str,
            sheet_name: str,
            data_range: str,
            categories_range: str | None = None,
            title: str | None = None,
            show_values: bool = True,
            *,
            target_sheet_name: str | None = None,
            position_cell: str = 'E5',
            x_axis_title: str | None = None,
            y_axis_title: str | None = None
    ) -> None:

        # Wybieramy rodzaj wykresu
        chart_class = {
            'bar': BarChart,
            'line': LineChart,
            'pie': PieChart
        }.get(chart_type.lower())

        if not chart_class:
            raise ValueError(f'Unsupported chart type: {chart_type}')

        chart = chart_class()
        if title:
            chart.title = title

        ws = self.get_sheet(sheet_name)

        data = Reference(
            worksheet=ws,
            range_string=f'{sheet_name}!{data_range}'
        )
        chart.add_data(data, titles_from_data=chart_type.lower != 'pie')

        # Dodaj etykiety
        if categories_range:
            cats = Reference(
                worksheet=ws,
                range_string=f'{sheet_name}!{categories_range}'
            )
            chart.set_categories(cats)

        if hasattr(chart, 'x_axis') and x_axis_title:
            chart.x_axis.title = x_axis_title
            chart.x_axis.title_font = Font(bold=True, size=16)

        if hasattr(chart, 'y_axis') and y_axis_title:
            chart.y_axis.title = y_axis_title
            chart.y_axis.title_font = Font(bold=True, size=16)
            chart.y_axis.majorGridlines = None

        if show_values and hasattr(chart, 'dataLabels'):
            chart.dataLabels = DataLabelList()
            chart.dataLabels.showVal = True

        chart.legend.position = 'b'
        # Inne opcje to: 't' (top), 'r' (right), 'l' (left), 'tr' (top right corner)
        chart.legend.overlay = False  # Sprawia, że legenda nie nakłada się na sam wykres

        if isinstance(chart, BarChart):
            chart.gapWidth = 120  # Zmniejsza odstęp między słupkami (domyślnie 150)
            # Mniejsze wartości = szersze słupki
            chart.overlap = 0  # Ustawia brak nakładania się słupków (dla wielu serii danych)

        # Wykres wstawiasz do tego samego arkusza
        if not target_sheet_name:
            ws.add_chart(chart, position_cell)

        # Wykres wstawiasz do nowego arkusza
        else:
            if target_sheet_name in self.workbook.sheetnames:
                del self.workbook[target_sheet_name]
            new_ws = self.workbook.create_sheet(title=target_sheet_name)
            new_ws.add_chart(chart, position_cell)

    # ------------------------------------------------------------------------------------------------------------------
    # METODY PRYWATNE
    # ------------------------------------------------------------------------------------------------------------------
    def _load_or_create(self) -> None:
        path = Path(self.filepath)
        if path.exists():
            self.workbook = load_workbook(self.filepath, data_only=True)
        else:
            self.workbook = Workbook()
            if self.workbook.active is not None:
                self.workbook.remove(self.workbook.active)

