from collections import defaultdict
from typing import Any, Callable

from openpyxl.styles import Font, PatternFill, Alignment

from invoice_group_excel.config import CFG
from invoice_group_excel.kmexcel.excel.manager import ExcelManager
from invoice_group_excel.kmexcel.excel.types.formula import FormulaTemplate
from openpyxl.utils import get_column_letter, column_index_from_string


class InvoiceAggregator(ExcelManager):
    def __init__(self, input_data_file: str, input_data_sheets: list[str]) -> None:
        super().__init__(input_data_file)
        self.input_data_sheets = input_data_sheets
        self._input_data: dict[str, list[dict[str, Any]]] = {}
        self._source_input_data: dict[str, list[dict[str, Any]]] = {}
        for data_sheet in input_data_sheets:
            self._input_data[data_sheet] = self.read_sheet(data_sheet,{CFG.summary_col_description[0]:CFG.summary_col_description[1]})
            self.create_source_sheet(data_sheet)

    def update_input_data(self) -> None:
        for sheet in self.get_sheet_names():
            self._input_data[sheet] = self.read_sheet(sheet,{CFG.summary_col_description[0]:CFG.summary_col_description[1]})

    def create_source_sheet(self,source_sheet_name: str) -> None:

        value_sheet = self._input_data[source_sheet_name]

        self.add_sheet("SOURCE."+source_sheet_name,value_sheet,formula_templates=self.summary_formula_templates(value_sheet))

    def get_input_data(self) -> dict[str,list[dict[str, Any]]]:
        return self._input_data

    def get_sheet_input_data(self,sheet_name: str) -> list[dict[str, Any]]:
        return self._input_data.get(sheet_name)

    def get_sheet_input_data_with_conditions(self,sheet_name: str, condition_fn: Callable[[dict[str,Any]],bool]) -> list[dict[str, Any]]:
        return [output for output in self.get_sheet_input_data(sheet_name) if condition_fn(output) ]
            #for key,value in output.items() if condition_fn(key,value)]

    def group_sheets_by_keys(self, source_sheet_name: str | None, keys_to_group: list[str]) -> dict[str, list[dict[str, Any]]]:
        if not source_sheet_name:
            source_sheet_name = self.input_data_sheets[0]
        output: dict[str, list[dict[str, Any]]] = {}
        input_data = self.get_sheet_input_data(source_sheet_name)

        # set_of_values: set[str] = set()
        # for row in input_data:
        #     set_of_values.add('.'.join([row[key] for key in keys_to_group]))#[row['Task'],row['TaskV2']]) )
        #
        # for value in set_of_values:
        #     output[value] = self.get_sheet_input_data_with_conditions(source_sheet_name,
        #         lambda x: all(x.get(key_to_group) == v for key_to_group,v in zip(keys_to_group,value.split('.'))))
        for row in input_data:
            group_key_tuple = tuple(row.get(k) for k in keys_to_group)
            group_key_str = "|".join(str(x) for x in group_key_tuple)  # tylko do nazwy klucza dict

            if group_key_str not in output:
                output[group_key_str] = []

            output[group_key_str].append(row)


        return output


    def add_sheets_grouped_by_key(self, source_sheet_name: str | None = None,*, key_to_group: list[str], sort_by: str | list[str] | None = None) -> None:
        if not source_sheet_name:
            source_sheet_name = self.input_data_sheets[0]

        sheets_grouped = self.group_sheets_by_keys(source_sheet_name, key_to_group)


        for key_sheet, value_sheet in sheets_grouped.items():
            if sort_by:
                if isinstance(sort_by, str):
                    sort_keys = [sort_by]
                else:
                    sort_keys = sort_by
                value_sheet.sort(key=lambda x: tuple(x[k] for k in sort_keys))
            self.add_sheet(key_sheet, value_sheet, formula_templates=self.summary_formula_templates(value_sheet))

    def add_sheet_with_conditions(self,
                                  source_sheet_name: str | None =None,*,
                                  output_sheet_name: str,
                                  condition_fn: Callable[[dict[str,Any]],bool],
                                  sort_by: str | list[str] | None = None,
                                  ) -> None:
        if not source_sheet_name:
            source_sheet_name = self.input_data_sheets[0]

        data_to_add = self.get_sheet_input_data_with_conditions(source_sheet_name, condition_fn)
        # zamiana jednego klucza w listę
        if sort_by:
            if isinstance(sort_by, str):
                sort_keys = [sort_by]
            else:
                sort_keys = sort_by
            data_to_add.sort( key=lambda row: tuple(row.get(k) for k in sort_keys))

        self.add_sheet(output_sheet_name,data_to_add,formula_templates=self.summary_formula_templates(data_to_add))


    def save_and_style(self,filepath: str | None = None) -> None:

        for key_sheet in self.get_sheet_names():
            self.style_all(
                key_sheet,
                style={
                    "font": Font(name="Calibri", size=11),
                    "alignment": Alignment(horizontal="center"),
                    "border_sides": {
                        "top": {"style": "thin", "color": "999999"},
                        "bottom": {"style": "thin", "color": "999999"},
                        "left": {"style": "thin", "color": "999999"},
                        "right": {"style": "thin", "color": "999999"},
                    },
                },
            )
            self.style_row(
                sheet_name=key_sheet,
                row_index=1,
                start_col_letter="A",
                style={
                    "font": Font(name="Calibri", size=12, bold=True, color="FFFFFF"),
                    "fill": PatternFill(fill_type="solid", fgColor="4F81BD"),
                    "alignment": Alignment(horizontal="center", vertical="center"),
                },
            )

            sheet = self.get_sheet(key_sheet)
            if sheet.cell(row=sheet.max_row, column=column_index_from_string(CFG.summary_col_description[0])).value == CFG.summary_col_description[1]:
                self.style_row(
                    sheet_name=key_sheet,
                    row_index=-1,
                    start_col_letter="A",
                    style={
                        "font": Font(name="Calibri", size=12, bold=True),# color="000000")
                        # "fill": PatternFill(fill_type="solid", fgColor="FFFFFF"),
                        # "alignment": Alignment(horizontal="center", vertical="center"),
                        "border_sides": {
                            "top": {"style": "medium", "color": "999999"},
                            "bottom": {"style": "medium", "color": "999999"},
                            "left": {"style": "medium", "color": "999999"},
                            "right": {"style": "medium", "color": "999999"},
                        },
                    },
                )

            for date_column in CFG.date_columns:
                self.set_column_format(key_sheet, date_column, CFG.date_format)
            self.autofit_column_widths(key_sheet, offset_dim=2)
        if filepath:
            for input_sheet in self.input_data_sheets:
                self.remove_sheet(input_sheet)
        self.save(filepath=filepath)

    # @staticmethod
    # def default_formula_templates(value_sheet:  list[dict[str, Any]]) -> list[FormulaTemplate]:
    #     formula_templates: list[FormulaTemplate] = []
    #
    #     for key, value in CFG.summary_description.items():
    #         formula_templates.append({"column": key, "start_row": len(value_sheet) + 2, "end_row": len(value_sheet) + 2,
    #          "template": value})
    #     return formula_templates

    @staticmethod
    def summary_formula_templates(value_sheet: list[dict[str, Any]]) -> list[FormulaTemplate]:
        formula_templates: list[FormulaTemplate] = []

        column_values: defaultdict = defaultdict(list)
        for item in value_sheet:
            for i, value in enumerate(item.values()):
                column_values[get_column_letter(i+1)].append(value)

        summary_ = {}
        for column in column_values.keys():
            if all(InvoiceAggregator.is_numeric_str(str(item)) for item in column_values[column]):
                summary_[column] = f"=SUM({column}2:{column}"+"{last})"
        # for key, value in CFG.summary_description.items():

        if len(summary_) > 0:
            for k,v in summary_.items():

                formula_templates.append(
                    {"column": CFG.summary_col_description[0], "start_row": len(value_sheet) + 2, "end_row": len(value_sheet) + 2,
                     "template": CFG.summary_col_description[1]})

                formula_templates.append({"column": k, "start_row": len(value_sheet) + 2, "end_row": len(value_sheet) + 2,
                                      "template": v})
        return formula_templates

    @staticmethod
    def is_numeric_str(x: str) -> bool:
        x = x.replace(",", ".")
        return x.replace(".", "", 1).isdigit()
