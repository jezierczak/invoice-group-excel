from typing import TypedDict
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import date
from invoice_group_excel.kmexcel.excel.manager import ExcelManager
from invoice_group_excel.kmexcel.excel.types.style import CellStyle, SheetStyle
from invoice_group_excel.kmexcel.excel.types.formula import FormulaDefinition, FormulaTemplate


class Person(TypedDict):
    name: str
    age: int
    department: str


class Event(TypedDict):
    name: str
    date: date


# def main() -> None:
#     data: list[Person] = [
#         {"name": "John", "age": 30, "department": "IT"},
#         {"name": "Anna", "age": 25, "department": "HR"},
#         {"name": "Peter", "age": 45, "department": "Finance"}
#     ]
#
#     data2: list[Event] = [
#         {"name": "Event A", "date": date(2025,1,5)},
#         {"name": "Event B", "date": date(2025,1,6)}
#     ]
#
#     formulas: list[FormulaDefinition] = [
#         {'cell': 'D2', 'formula': '=IF(B2>30, "Senior", "Junior")'},
#         {'cell': 'D3', 'formula': '=IF(B3>30, "Senior", "Junior")'}
#     ]
#
#     formulas2: list[FormulaDefinition] = [
#         {'cell': 'C2', 'formula': '=YEAR(B2)'}
#     ]
#
#     formula_templates: list[FormulaTemplate] = [
#         {
#             'column': 'E',
#             'template': '="Dept: " & C{row}'
#         },
#         {
#             'column': 'F',
#             'start_row': len(data) + 2,
#             'end_row': len(data) + 2,
#             'template': '=SUM(B2:B{last})'
#         },
#     ]
#
#     excel = ExcelManager('people.xlsx')
#     excel.add_sheet('People1', data, formulas=formulas, formula_templates=formula_templates)
#     excel.add_sheet('People2', data2, formulas=formulas2)
#
#     # excel.add_column('People1', values=[111, 222, 333], col_letter='A', row_start=1)
#     # excel.add_row('People1', data={'name': 'John', 'age': 30, 'department': 'IT'}, row_index=10, col_letter='D')
#     # excel.add_cell('People1', 'G9', 'xxx')
#
#     excel.style_all('People1', style={
#         'font': Font(name='Verdana', size=12),
#         'alignment': Alignment(horizontal='center'),
#         'border_sides': {
#             'top': {'style': 'thick', 'color': 'AABB11'},
#             'bottom': {'style': 'thick', 'color': 'AABB11'},
#             'left': {'style': 'thick', 'color': 'AABB11'},
#             'right': {'style': 'thick', 'color': 'AABB11'}
#         }
#     })
#
#     excel.style_row(
#         sheet_name='People1',
#         row_index=1,
#         start_col_letter='B',
#         style={
#         'font': Font(name='Verdana', size=12, bold=True, color='FFFFFF'),
#         'fill': PatternFill(fill_type='solid', fgColor='AA8712'),
#         'alignment': Alignment(horizontal='center', vertical='center')
#     })
#
#     excel.style_column(
#         sheet_name='People1',
#         col_letter='A',
#         start_row=2,
#         style={
#         'font': Font(name='Verdana', size=12, bold=True, color='FFFFFF'),
#         'fill': PatternFill(fill_type='solid', fgColor='4AA1BD'),
#         'alignment': Alignment(horizontal='center', vertical='center')
#     })
#
#     excel.style_cell('People1', 'F5', style={
#         'font': Font(name='Verdana', size=22, bold=True, color='CCBB54'),
#         'alignment': Alignment(horizontal='right', vertical='center')
#     })
#
#
#     excel.set_column_format('People1', 'B', '0.00')
#     excel.set_column_format('People2', 'B', 'MM-DD-YYYY')
#     excel.autofit_column_widths('People1')
#
#     excel.add_new_chart(
#         chart_type='bar',
#         sheet_name='People1',
#         data_range='B2:B4',
#         categories_range='A2:A4',
#         title='People age',
#         position_cell='H2',
#         show_values=False,
#         x_axis_title='X',
#         y_axis_title='Y'
#     )
#
#     excel.add_new_chart(
#         chart_type='line',
#         sheet_name='People1',
#         data_range='B2:B4',
#         categories_range='A2:A4',
#         title='Age',
#         target_sheet_name='PeopleChart',
#         position_cell='A1',
#         show_values=False,
#         x_axis_title='X',
#         y_axis_title='Y'
#     )
#
#     excel.remove_sheet('PeopleChart')
#     excel.save()


def main() -> None:
    data = [
        {"name": "John", "age": 30, "department": "IT"},
        {"name": "Anna", "age": 25, "department": "HR"},
        {"name": "Peter", "age": 45, "department": "Finance"},
    ]


    data2 = [
        {"name": "Event A", "date": date(2025, 1, 5)},
        {"name": "Event B", "date": date(2025, 1, 6)},
    ]

    formulas = [
        {"cell": "D2", "formula": '=IF(B2>30, "Senior", "Junior")'},
        {"cell": "D3", "formula": '=IF(B3>30, "Senior", "Junior")'},
    ]

    formulas2 = [
        {"cell": "C2", "formula": "=YEAR(B2)"}
    ]

    formula_templates = [
        {"column": "E", "template": '="Dept: " & C{row}'},
        {"column": "F", "start_row": len(data) + 2, "end_row": len(data) + 2, "template": "=SUM(B2:B{last})"},
    ]

    excel = ExcelManager("km.xlsx")
    excel.add_sheet("People1", data, formulas=formulas, formula_templates=formula_templates)
    excel.add_sheet("People2", data2, formulas=formulas2)

    excel.style_all(
        "People1",
        style={
            "font": Font(name="Calibri", size=11),
            "alignment": Alignment(horizontal="center"),
            "border_sides": {
                "top": {"style": "thin", "color": "999999"},
                "bottom": {"style": "thin", "color": "999999"},
                "left": {"style": "thin", "color": "999999"},
                "right": {"style": "thin", "color": "999999"},
            },
        },
    )

    excel.style_row(
        sheet_name="People1",
        row_index=1,
        start_col_letter="A",
        style={
            "font": Font(name="Calibri", size=12, bold=True, color="FFFFFF"),
            "fill": PatternFill(fill_type="solid", fgColor="4F81BD"),
            "alignment": Alignment(horizontal="center", vertical="center"),
        },
    )

    excel.style_column(
        sheet_name="People1",
        col_letter="A",
        start_row=2,
        style={
            "font": Font(name="Calibri", size=12, bold=True, color="2F4F4F"),
            "fill": PatternFill(fill_type="solid", fgColor="D9E1F2"),
            "alignment": Alignment(horizontal="left"),
        },
    )

    excel.style_cell(
        "People1",
        "F5",
        style={
            "font": Font(name="Calibri", size=16, bold=True, color="FF6600"),
            "alignment": Alignment(horizontal="right", vertical="center"),
        },
    )

    excel.set_column_format("People1", "B", "0.00")
    excel.set_column_format("People2", "B", "MM-DD-YYYY")

    excel.autofit_column_widths("People1", offset_dim=2)

    # -------------------------------
    # Wykres 1: BarChart w tym samym arkuszu
    # -------------------------------
    excel.add_new_chart(
        chart_type="bar",
        sheet_name="People1",
        data_range="B2:B4",
        categories_range="A2:A4",
        title="Wiek pracowników",
        position_cell="H2",
        show_values=True,
        x_axis_title="Pracownik",
        y_axis_title="Wiek",
    )

    # -------------------------------
    # Wykres 2: LineChart w osobnym arkuszu
    # -------------------------------
    excel.add_new_chart(
        chart_type="line",
        sheet_name="People1",
        data_range="B2:B4",
        categories_range="A2:A4",
        title="Trend wieku",
        target_sheet_name="PeopleChart",
        position_cell="B3",
        show_values=True,
        x_axis_title="Imię",
        y_axis_title="Liczba lat",
    )

    # excel.remove_sheet("PeopleChart")
    excel.save()

if __name__ == '__main__':
    main()
